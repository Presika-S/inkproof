
import hashlib
import io
import uuid
from datetime import date

import docx
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, Field
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import models as m
from .db import get_db
from .security import get_current_user, require_reviewer

router = APIRouter(prefix="/api/evidence", tags=["Evidence"])

MAX_FILE_SIZE = 5_000_000


class ProposalIn(BaseModel):
    document_id: str
    source_quote: str = Field(min_length=5, max_length=5000)
    fact_text: str = Field(min_length=5, max_length=300)
    keywords: list[str] = Field(default_factory=list)
    value: float
    unit: str = Field(min_length=1, max_length=20)
    period: float | None = None
    period_unit: str | None = None
    cumulative: bool = False
    rank: bool = False
    valid_until: str
    owner: str = "Research team"


class RejectIn(BaseModel):
    note: str = Field(min_length=3, max_length=300)


def extract_text(filename: str, data: bytes) -> str:
    name = filename.lower()

    try:
        if name.endswith(".txt"):
            text = data.decode("utf-8-sig")
        elif name.endswith(".docx"):
            document = docx.Document(io.BytesIO(data))
            text = "\n".join(p.text for p in document.paragraphs)
        elif name.endswith(".pdf"):
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise HTTPException(
                    400, "Password-protected PDFs are unsupported"
                )
            text = "\n".join(
                page.extract_text() or "" for page in reader.pages
            )
            if not text.strip():
                raise HTTPException(
                    400, "No selectable PDF text found; OCR is needed"
                )
        else:
            raise HTTPException(400, "Upload a PDF, DOCX or TXT file")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Could not read this document")

    text = text.strip()
    if not text:
        raise HTTPException(400, "The document contains no text")

    return text


def document_dict(document):
    return {
        "id": document.id,
        "filename": document.filename,
        "version": document.version,
        "status": document.status,
        "characters_extracted": len(document.extracted_text),
        "created_at": document.created_at.isoformat(),
    }


def proposal_dict(proposal):
    return {
        "id": proposal.id,
        "document_id": proposal.document_id,
        "source_quote": proposal.source_quote,
        "fact_text": proposal.fact_text,
        "keywords": proposal.keywords,
        "value": proposal.value,
        "unit": proposal.unit,
        "period": proposal.period,
        "period_unit": proposal.period_unit,
        "cumulative": proposal.cumulative,
        "rank": proposal.rank,
        "valid_until": proposal.valid_until,
        "owner": proposal.owner,
        "status": proposal.status,
        "created_by": proposal.created_by,
        "reviewer_id": proposal.reviewer_id,
        "reviewed_at": (
            proposal.reviewed_at.isoformat()
            if proposal.reviewed_at else None
        ),
        "fact_id": proposal.fact_id,
        "created_at": proposal.created_at.isoformat(),
    }


def log_event(db, event_type, user_id, note, before=None, after=None):
    db.add(
        m.Event(
            type=event_type,
            note=note[:300],
            before_value=before,
            after_value=after,
            user_id=user_id,
        )
    )


@router.post("/documents")
async def upload_document(
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    filename = (file.filename or "").strip()
    if not filename:
        raise HTTPException(400, "Filename is required")

    data = await file.read(MAX_FILE_SIZE + 1)
    if len(data) > MAX_FILE_SIZE:
        raise HTTPException(413, "File too large (5 MB max)")

    text = extract_text(filename, data)
    digest = hashlib.sha256(data).hexdigest()

    existing = list(
        db.scalars(
            select(m.EvidenceDocument).where(
                m.EvidenceDocument.filename == filename
            )
        )
    )
    version = max((d.version for d in existing), default=0) + 1

    for old_document in existing:
        if old_document.status == "active":
            old_document.status = "superseded"

            # Existing fact links must no longer be treated as current.
            old_links = db.scalars(
                select(m.FactEvidence).where(
                    m.FactEvidence.document_id == old_document.id
                )
            )
            for link in old_links:
                link.stale = True

    document = m.EvidenceDocument(
        id=str(uuid.uuid4()),
        filename=filename,
        content_hash=digest,
        extracted_text=text,
        version=version,
        status="active",
        uploaded_by=user.id,
    )

    db.add(document)
    db.commit()
    db.refresh(document)

    return {
        **document_dict(document),
        "message": "Document uploaded successfully",
    }


@router.get("/documents")
def list_documents(
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    documents = db.scalars(
        select(m.EvidenceDocument).order_by(
            m.EvidenceDocument.created_at.desc()
        )
    )
    return [document_dict(d) for d in documents]


@router.get("/documents/{document_id}")
def get_document(
    document_id: str,
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    document = db.get(m.EvidenceDocument, document_id)
    if not document:
        raise HTTPException(404, "Document not found")

    return {
        **document_dict(document),
        "text": document.extracted_text,
    }


@router.post("/proposals")
def create_proposal(
    payload: ProposalIn,
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    document = db.get(m.EvidenceDocument, payload.document_id)
    if not document:
        raise HTTPException(404, "Document not found")

    if document.status != "active":
        raise HTTPException(
            409, "Cannot create a proposal from a superseded document"
        )

    if payload.source_quote not in document.extracted_text:
        raise HTTPException(
            400, "The exact source quote must appear in the document"
        )

    try:
        date.fromisoformat(payload.valid_until)
    except ValueError:
        raise HTTPException(
            400, "valid_until must use YYYY-MM-DD format"
        )

    if not payload.keywords:
        raise HTTPException(400, "At least one keyword is required")

    proposal = m.EvidenceProposal(
        id=str(uuid.uuid4()),
        document_id=document.id,
        source_quote=payload.source_quote,
        fact_text=payload.fact_text,
        keywords=payload.keywords,
        value=payload.value,
        unit=payload.unit,
        period=payload.period,
        period_unit=payload.period_unit,
        cumulative=payload.cumulative,
        rank=payload.rank,
        valid_until=payload.valid_until,
        owner=payload.owner,
        status="pending",
        created_by=user.id,
    )

    db.add(proposal)
    log_event(
        db,
        "propose",
        user.id,
        f"Evidence fact proposal created: {proposal.id}",
        after=payload.fact_text,
    )
    db.commit()
    db.refresh(proposal)

    return proposal_dict(proposal)


@router.get("/proposals")
def list_proposals(
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    proposals = db.scalars(
        select(m.EvidenceProposal).order_by(
            m.EvidenceProposal.created_at.desc()
        )
    )
    return [proposal_dict(p) for p in proposals]


@router.get("/proposals/pending")
def list_pending_proposals(
    db: Session = Depends(get_db),
    reviewer: m.User = Depends(require_reviewer),
):
    proposals = db.scalars(
        select(m.EvidenceProposal)
        .where(m.EvidenceProposal.status == "pending")
        .order_by(m.EvidenceProposal.created_at.asc())
    )
    return [proposal_dict(p) for p in proposals]


@router.post("/proposals/{proposal_id}/approve")
def approve_proposal(
    proposal_id: str,
    db: Session = Depends(get_db),
    reviewer: m.User = Depends(require_reviewer),
):
    proposal = db.get(m.EvidenceProposal, proposal_id)
    if not proposal:
        raise HTTPException(404, "Proposal not found")

    if proposal.status != "pending":
        raise HTTPException(409, "This proposal has already been reviewed")

    if proposal.created_by == reviewer.id:
        raise HTTPException(
            403, "You cannot approve your own evidence proposal"
        )

    document = db.get(m.EvidenceDocument, proposal.document_id)
    if not document or document.status != "active":
        raise HTTPException(
            409, "The source document is no longer active"
        )

    if proposal.source_quote not in document.extracted_text:
        raise HTTPException(
            409, "The source quote is no longer present in the document"
        )

    fact_id = f"F-{uuid.uuid4().hex[:12].upper()}"

    fact = m.Fact(
        id=fact_id,
        text=proposal.fact_text,
        keywords=proposal.keywords,
        value=proposal.value,
        unit=proposal.unit,
        period=proposal.period,
        period_unit=proposal.period_unit,
        cumulative=proposal.cumulative,
        rank=proposal.rank,
        valid_until=proposal.valid_until,
        owner=proposal.owner,
        version=1,
        approved=True,
    )

    evidence_link = m.FactEvidence(
        id=str(uuid.uuid4()),
        fact_id=fact_id,
        document_id=document.id,
        source_quote=proposal.source_quote,
        source_version=document.version,
        stale=False,
    )

    proposal.status = "approved"
    proposal.reviewer_id = reviewer.id
    proposal.reviewed_at = m.now()
    proposal.fact_id = fact_id

    db.add(fact)
    db.add(evidence_link)
    log_event(
        db,
        "approve",
        reviewer.id,
        f"Evidence proposal approved: {proposal.id}",
        before="pending",
        after=f"approved; fact_id={fact_id}",
    )

    db.commit()

    return {
        "message": "Proposal approved and fact created",
        "proposal_id": proposal.id,
        "fact_id": fact_id,
        "document_id": document.id,
        "source_quote": proposal.source_quote,
        "document_version": document.version,
    }


@router.post("/proposals/{proposal_id}/reject")
def reject_proposal(
    proposal_id: str,
    payload: RejectIn,
    db: Session = Depends(get_db),
    reviewer: m.User = Depends(require_reviewer),
):
    proposal = db.get(m.EvidenceProposal, proposal_id)
    if not proposal:
        raise HTTPException(404, "Proposal not found")

    if proposal.status != "pending":
        raise HTTPException(409, "This proposal has already been reviewed")

    if proposal.created_by == reviewer.id:
        raise HTTPException(
            403, "You cannot reject your own evidence proposal"
        )

    proposal.status = "rejected"
    proposal.reviewer_id = reviewer.id
    proposal.reviewed_at = m.now()

    log_event(
        db,
        "reject",
        reviewer.id,
        f"Evidence proposal rejected: {payload.note}",
        before="pending",
        after="rejected",
    )

    db.commit()

    return {
        "message": "Proposal rejected",
        "proposal_id": proposal.id,
        "status": proposal.status,
    }


@router.get("/facts/{fact_id}/sources")
def get_fact_sources(
    fact_id: str,
    db: Session = Depends(get_db),
    user: m.User = Depends(get_current_user),
):
    fact = db.get(m.Fact, fact_id)
    if not fact:
        raise HTTPException(404, "Fact not found")

    links = db.scalars(
        select(m.FactEvidence).where(
            m.FactEvidence.fact_id == fact_id
        )
    )

    results = []
    for link in links:
        document = db.get(m.EvidenceDocument, link.document_id)
        results.append({
            "document_id": link.document_id,
            "filename": document.filename if document else None,
            "source_quote": link.source_quote,
            "document_version": link.source_version,
            "current_document_version": (
                document.version if document else None
            ),
            "stale": link.stale,
        })

    return results
